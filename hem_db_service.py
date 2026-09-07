"""Koneksi & insert data "HEM / Input Data Semesta" ke Postgres (addon
Railway) -- dipakai oleh endpoint /api/hem/insert di app.py.

Railway OTOMATIS nyuntik env var DATABASE_URL begitu addon Postgres
di-attach ke service ini (lihat config.DATABASE_URL), jadi tidak perlu
diisi manual.

PENTING -- HEM_FIELDS di bawah SENGAJA disalin persis dari COLUMN_GROUPS
di templates/hem.html (key & urutannya harus SAMA PERSIS, itu yang
dipakai frontend buat kirim payload JSON ke endpoint ini). Kalau nanti
nambah/hapus/ubah nama kolom di salah satu sisi, sisi yang satu lagi
WAJIB diupdate juga supaya tetap sinkron -- tidak ada mekanisme
otomatis yang menjaga keduanya tetap sama.
"""
import datetime
import decimal
import re

import psycopg2
import psycopg2.extras

import config

TABLE_NAME = "data_semesta"

# (key, type) -- type: "number" | "text" | "textarea" | "date" | "datetime"
HEM_FIELDS = [
    ("no", "number"), ("tahun", "number"), ("no_order", "text"),
    ("ihld_lop_id", "text"), ("nama_proyek", "text"), ("status_ihld", "text"),
    ("tipe_desain", "text"), ("prioritas", "text"),

    ("status_wo_tif", "text"), ("new_region_ta", "text"), ("region_tif", "text"),
    ("reg_lama", "text"), ("witel_lama", "text"), ("branch", "text"),
    ("sto", "text"), ("sc", "text"), ("datek", "text"), ("tipe_deploy_actual", "text"),

    ("layanan", "text"), ("nde_wo", "text"), ("tanggal_nde_wo", "date"),
    ("nde_permohonan_ut", "text"),

    ("panjang_kabel_meter", "number"), ("id_pr_material", "text"), ("pid", "text"),
    ("sap", "text"), ("pr", "text"), ("po", "text"),

    ("progress_w01", "text"), ("progress_w02", "text"),
    ("progress_jt_last_update", "text"), ("progress", "text"),
    ("keterangan_detail", "textarea"), ("umur_order", "number"),
    ("grouping_umur_order", "text"), ("issue", "textarea"),

    ("target_fi", "date"), ("tanggal_fi", "date"), ("tgl_go_live", "date"),
    ("bulan_golive", "text"), ("tgl_ut", "date"), ("target_weekly_bast", "text"),

    ("status_drop", "text"), ("status_ut", "text"), ("status_rekon", "text"),
    ("bast", "text"), ("status_ba_drop", "text"), ("tanggal_ba_drop", "date"),
    ("ket_ba_drop", "textarea"), ("ba_redesign", "text"), ("lact", "text"),
    ("baut", "text"),

    ("total_boq_ihld", "number"), ("total_boq_wo", "number"), ("total_drm", "number"),
    ("total_boq_actual_ut_rekon", "number"), ("material_ut", "text"),
    ("jasa_ut", "text"), ("total_ut", "number"), ("nilai_perizinan", "number"),
    ("kenaikan", "number"), ("selisih_nilai_wo_ihld", "number"),
    ("persen_perubahan_nilai", "number"),

    ("nama_mitra", "text"), ("jumlah_manpower", "number"), ("no_wo_smile", "text"),
    ("nama_smile", "text"), ("persen_smile", "number"), ("status_smile", "text"),

    ("sp", "text"), ("nomor_sp", "text"),
]
HEM_FIELD_TYPE = dict(HEM_FIELDS)
HEM_FIELD_KEYS = [k for k, _ in HEM_FIELDS]

# Label tampilan per kolom -- DISALIN PERSIS dari LABEL_BY_KEY (COLUMN_GROUPS)
# di templates/hem.html, sama alasannya dengan HEM_FIELDS di atas: dipakai
# buat header export Excel & label kolom Detail di dashboard (hem_dashboard.html)
# supaya labelnya konsisten dengan form input, BUKAN dijaga otomatis --
# kalau salah satu sisi berubah, sisi yang lain wajib diupdate manual juga.
HEM_FIELD_LABELS = {
    "no": "No", "tahun": "Tahun", "no_order": "No Order",
    "ihld_lop_id": "iHLD LoP ID", "nama_proyek": "Nama Proyek",
    "status_ihld": "Status iHLD", "tipe_desain": "Tipe Desain", "prioritas": "Prioritas",

    "status_wo_tif": "Status WO (TIF)", "new_region_ta": "New Region TA",
    "region_tif": "Region TIF", "reg_lama": "Reg Lama", "witel_lama": "Witel Lama",
    "branch": "Branch", "sto": "STO", "sc": "SC", "datek": "Datek",
    "tipe_deploy_actual": "Tipe Deploy Actual",

    "layanan": "Layanan", "nde_wo": "NDE WO", "tanggal_nde_wo": "Tanggal NDE/WO",
    "nde_permohonan_ut": "NDE Permohonan UT",

    "panjang_kabel_meter": "Panjang Kabel (meter)", "id_pr_material": "ID PR Material",
    "pid": "PID", "sap": "SAP", "pr": "PR", "po": "PO",

    "progress_w01": "Progress W-01", "progress_w02": "Progress W-02",
    "progress_jt_last_update": "Progress JT Last Update", "progress": "Progress",
    "keterangan_detail": "Keterangan / Detail", "umur_order": "Umur Order",
    "grouping_umur_order": "Grouping Umur Order", "issue": "Issue",

    "target_fi": "Target FI", "tanggal_fi": "Tanggal FI", "tgl_go_live": "Tgl Go Live",
    "bulan_golive": "Bulan Golive", "tgl_ut": "Tgl UT",
    "target_weekly_bast": "Target Weekly BAST",

    "status_drop": "Status Drop", "status_ut": "Status UT", "status_rekon": "Status Rekon",
    "bast": "BAST", "status_ba_drop": "Status BA Drop", "tanggal_ba_drop": "Tanggal BA Drop",
    "ket_ba_drop": "Ket BA Drop", "ba_redesign": "BA Redesign", "lact": "LACT", "baut": "BAUT",

    "total_boq_ihld": "Total BOQ IHLD", "total_boq_wo": "Total BOQ WO",
    "total_drm": "Total DRM", "total_boq_actual_ut_rekon": "Total BOQ Actual UT - Rekon",
    "material_ut": "Material UT", "jasa_ut": "Jasa UT", "total_ut": "Total UT",
    "nilai_perizinan": "Nilai Perizinan", "kenaikan": "Kenaikan",
    "selisih_nilai_wo_ihld": "Selisih Nilai WO & IHLD",
    "persen_perubahan_nilai": "% Perub Nilai",

    "nama_mitra": "Nama Mitra", "jumlah_manpower": "Jumlah Manpower",
    "no_wo_smile": "No WO Smile", "nama_smile": "Nama Smile",
    "persen_smile": "% Smile", "status_smile": "Status Smile",

    "sp": "SP", "nomor_sp": "Nomor SP",
}

# Kolom yang dipakai sebagai filter/rekap di dashboard -- daftar nilai
# uniknya diambil langsung dari data yang ada (get_filter_options), BUKAN
# daftar tetap, supaya otomatis ikut kalau ada nilai baru masuk lewat
# form/upload di hem.html.
DASHBOARD_FILTER_COLS = ["new_region_ta", "branch", "status_ihld", "prioritas"]


def _row_value_to_json_safe(v):
    """Baris hasil query psycopg2 (RealDictCursor) bisa berisi
    decimal.Decimal (kolom NUMERIC) atau datetime.date/datetime (kolom
    DATE/TIMESTAMP) -- keduanya TIDAK bisa langsung di-jsonify Flask.
    Konversi ke float/ISO-string di sini, sekali, dipakai semua fungsi
    baca di bawah."""
    if v is None:
        return None
    if isinstance(v, decimal.Decimal):
        return float(v)
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    return v


def get_dashboard_rows(filters: dict = None):
    """Baca baris dari data_semesta buat dashboard. `filters` dict
    opsional {col: [nilai, ...]} -- HANYA kolom di DASHBOARD_FILTER_COLS
    yang diterima (kolom lain diabaikan diam-diam, jaga-jaga input dari
    luar). Filter diterapkan di level SQL (WHERE ... IN (...)) supaya
    tetap ringan walau baris di tabel banyak, bukan difilter belakangan
    di Python. Return: list of dict {"id": ..., **HEM_FIELD_KEYS}, semua
    value sudah JSON-safe (lihat _row_value_to_json_safe)."""
    conn = get_connection()
    try:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cols_sql = ", ".join(f'"{k}"' for k in HEM_FIELD_KEYS)
            where_clauses = []
            params = []
            if filters:
                for col in DASHBOARD_FILTER_COLS:
                    values = [v for v in (filters.get(col) or []) if v]
                    if values:
                        placeholders = ", ".join(["%s"] * len(values))
                        where_clauses.append(f'"{col}" IN ({placeholders})')
                        params.extend(values)
            where_sql = f"WHERE {' AND '.join(where_clauses)}" if where_clauses else ""
            cur.execute(f'SELECT id, {cols_sql} FROM {TABLE_NAME} {where_sql} ORDER BY id', params)
            raw_rows = cur.fetchall()
    finally:
        conn.close()

    return [
        {k: _row_value_to_json_safe(v) for k, v in dict(r).items()}
        for r in raw_rows
    ]


def get_filter_options():
    """Daftar nilai unik utk dropdown/checkbox filter dashboard (New
    Region TA, Branch, Status iHLD, Prioritas), diambil langsung dari isi
    tabel saat ini -- bukan daftar tetap yang bisa basi kalau ada nilai
    baru masuk lewat form manual/upload di hem.html."""
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            options = {}
            for col in DASHBOARD_FILTER_COLS:
                cur.execute(
                    f'SELECT DISTINCT "{col}" FROM {TABLE_NAME} '
                    f'WHERE "{col}" IS NOT NULL AND TRIM("{col}") <> \'\' '
                    f'ORDER BY 1'
                )
                options[col] = [r[0] for r in cur.fetchall()]
    finally:
        conn.close()
    return options


def _to_num(v):
    try:
        return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def get_dashboard_summary(rows: list):
    """Ringkasan KPI + rekap "New Region TA" x "Status iHLD" (baris x
    kolom, kayak tabel Rekap Regional/Branch di dashboard PT3), dihitung
    di Python dari `rows` yang SUDAH diambil get_dashboard_rows() (biar
    tidak query ulang -- 1x baca dipakai buat rows mentah, summary, DAN
    export sekaligus)."""
    total = len(rows)
    total_nilai_wo = sum(_to_num(r.get("total_boq_wo")) for r in rows)
    total_nilai_ihld = sum(_to_num(r.get("total_boq_ihld")) for r in rows)
    total_ut = sum(_to_num(r.get("total_ut")) for r in rows)
    golive_count = sum(1 for r in rows if (r.get("tgl_go_live") or "").strip())
    drop_count = sum(1 for r in rows if (r.get("status_drop") or "").strip())

    regions = {}       # {region: {"__total__": n, status1: n, status2: n, ...}}
    statuses_seen = []  # urutan kemunculan pertama, dipakai jadi kolom tabel rekap
    for r in rows:
        region = (r.get("new_region_ta") or "").strip() or "(Tanpa Region)"
        status = (r.get("status_ihld") or "").strip() or "(Tanpa Status)"
        if status not in statuses_seen:
            statuses_seen.append(status)
        bucket = regions.setdefault(region, {"__total__": 0})
        bucket[status] = bucket.get(status, 0) + 1
        bucket["__total__"] += 1

    return {
        "total_order": total,
        "total_nilai_wo": total_nilai_wo,
        "total_nilai_ihld": total_nilai_ihld,
        "total_ut": total_ut,
        "golive_count": golive_count,
        "drop_count": drop_count,
        "statuses": statuses_seen,
        "regions": regions,
    }


def build_export_workbook(filters: dict = None):
    """Export SEMUA kolom data_semesta (sesuai filter dashboard yang lagi
    aktif, sama seperti Export Data A-AP di dashboard PT3) sebagai
    .xlsx -- 1 baris = 1 record, header pakai HEM_FIELD_LABELS supaya
    lebih enak dibaca daripada nama kolom mentah (snake_case)."""
    import io
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    rows = get_dashboard_rows(filters)
    headers = ["ID"] + [HEM_FIELD_LABELS.get(k, k) for k in HEM_FIELD_KEYS]

    wb = Workbook()
    sheet = wb.active
    sheet.title = "Data Semesta"
    sheet.append(headers)
    sheet.freeze_panes = "A2"

    for r in rows:
        sheet.append([r.get("id", "")] + [r.get(k, "") if r.get(k) is not None else "" for k in HEM_FIELD_KEYS])

    for i in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(i)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def get_connection():
    if not config.DATABASE_URL:
        raise RuntimeError(
            "DATABASE_URL belum ke-set. Attach addon Postgres di Railway ke "
            "service ini dulu (Railway otomatis nyuntik env var DATABASE_URL "
            "begitu sudah di-attach, tidak perlu diisi manual)."
        )
    return psycopg2.connect(config.DATABASE_URL)


def _cast_value(field_type, raw, key=None):
    """String dari JS -> tipe Python yang cocok buat psycopg2 ("" / None
    jadi NULL). number di-parse lewat _parse_number() (kecuali kolom
    "tahun", lihat _extract_year_number()); date/datetime di-parse &
    dinormalisasi ke ISO lewat _parse_date()/_parse_datetime(); text/
    textarea dikirim apa adanya sebagai string."""
    if raw is None:
        return None
    val = str(raw).strip()
    if val == "":
        return None
    if field_type == "number":
        if key == "tahun":
            return _extract_year_number(val)
        return _parse_number(val)
    if field_type == "date":
        return _parse_date(val)
    if field_type == "datetime":
        return _parse_datetime(val)
    return val


# Kolom "tahun" di file sumber ternyata bercampur: angka murni ("2026")
# ATAU label + tahun ("CO 2025", kemungkinan singkatan "Carry Over 2025").
# Atas keputusan pengguna, label seperti "CO" dibuang -- hanya 4 digit
# tahunnya yang disimpan (mis. "CO 2025" -> 2025). Info "CO" itu sendiri
# TIDAK disimpan di kolom manapun setelah ini.
_YEAR_RE = re.compile(r"(\d{4})")


def _extract_year_number(val):
    m = _YEAR_RE.search(val)
    if not m:
        raise ValueError(f"'{val}' tidak mengandung tahun (4 digit angka) yang valid")
    return int(m.group(1))


# Token error formula Excel & placeholder "kosong" -- atas keputusan
# pengguna, ini dianggap NULL (baris tetap masuk, kolom itu kosong)
# alih-alih menggagalkan baris.
_NUMBER_NULL_TOKENS = {"#n/a", "#ref!", "#value!", "#div/0!", "#null!", "#num!", "#name?", "-"}

# Format Indonesia: titik = pemisah ribuan, koma = desimal, mis.
# "7.111.783,00" atau "-4.609.953". Regex mengharuskan grup 3 digit
# persis setelah tiap titik supaya tidak salah anggap titik desimal
# biasa (mis. "12.5") sebagai format Indonesia.
_ID_THOUSANDS_RE = re.compile(r'^-?\d{1,3}(\.\d{3})+(,\d+)?$')
_PLAIN_COMMA_DECIMAL_RE = re.compile(r'^-?\d+,\d+$')


def _parse_number(val):
    """Normalisasi angka dari file Excel/CSV: buang tanda '%' (nilai
    persen disimpan sebagai angka utuh, bukan pecahan -- keputusan
    pengguna), pahami format ribuan/desimal ala Indonesia, dan anggap
    token error Excel / '-' sebagai NULL. Sisanya di-parse lewat
    float() seperti biasa."""
    v = val.strip()
    if v.lower() in _NUMBER_NULL_TOKENS:
        return None
    if v.endswith("%"):
        v = v[:-1].strip()
    if _ID_THOUSANDS_RE.match(v):
        v = v.replace(".", "").replace(",", ".")
    elif _PLAIN_COMMA_DECIMAL_RE.match(v):
        v = v.replace(",", ".")
    try:
        f = float(v)
    except ValueError:
        raise ValueError(f"'{val}' bukan angka yang valid")
    return int(f) if f.is_integer() else f


# Format tanggal yang diterima, dicoba berurutan sampai salah satu cocok.
# "%Y-%m-%d" -> ISO, dikirim oleh <input type="date"> di form manual.
# "%d/%m/%Y" dan "%d-%m-%Y" -> format Indonesia umum dari file Excel/CSV
# (mis. "21/07/2026"). Pola "DD-Mon-YY" (mis. "05-Agu-26") ditangani
# terpisah lewat _MONTH_ABBR karena singkatan bulannya campur
# Indonesia/Inggris (Agu/Aug, Mei/May, Okt/Oct, Des/Dec, dst).
_DATE_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"]

_DATETIME_FORMATS = [
    "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M",   # dari <input type="datetime-local">
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M",
    "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M",
]

_MONTH_ABBR = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4,
    "mei": 5, "may": 5,
    "jun": 6, "jul": 7,
    "agu": 8, "aug": 8,
    "sep": 9,
    "okt": 10, "oct": 10,
    "nov": 11,
    "des": 12, "dec": 12,
}
_MONTH_ABBR_DATE_RE = re.compile(r'^(\d{1,2})-([A-Za-z]{3})-(\d{2,4})$')

# Serial date Excel: hari sejak 1899-12-30 (termasuk bug tahun kabisat
# 1900 bawaan Excel/Lotus). Dipakai kalau sel Excel bertipe tanggal asli
# tapi kebetulan terbaca sebagai angka mentah, bukan teks terformat.
_EXCEL_EPOCH = datetime.date(1899, 12, 30)

# "DROP" di kolom target_fi/tanggal_fi/tgl_go_live artinya target itu
# sendiri dibatalkan (bukan data rusak) -- atas keputusan pengguna,
# dianggap NULL, sisanya (mayoritas tanggal asli di kolom2 itu) tetap
# tersimpan sebagai tanggal. Data tanggal yang benar-benar rusak/typo
# (mis. "170Feb026") SENGAJA tidak ditangkap di sini -- tetap gagal
# per baris supaya dikoreksi manual di file sumbernya.
_DATE_NULL_TOKENS = {"drop"}


def _parse_date(val):
    """Normalisasi string tanggal ke ISO 'YYYY-MM-DD'.

    PENTING: kalau string mentah (mis. '21/07/2026') dikirim apa adanya
    ke Postgres tanpa dinormalisasi, Postgres membacanya pakai DateStyle
    default (MDY -- bulan/hari/tahun), jadi '21' dibaca sebagai bulan
    dan meledak dengan error "date/time field value out of range" untuk
    tanggal yang harinya > 12. Di sinilah tempatnya divalidasi & diubah
    ke ISO SEBELUM sampai ke Postgres -- dan karena ini dipanggil per
    baris (lewat _cast_value di dalam loop insert_rows), satu tanggal
    salah format hanya menggagalkan baris itu, bukan seluruh batch."""
    if val.lower() in _DATE_NULL_TOKENS:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.datetime.strptime(val, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    m = _MONTH_ABBR_DATE_RE.match(val)
    if m:
        day_s, mon_s, year_s = m.groups()
        mon = _MONTH_ABBR.get(mon_s.lower())
        if mon:
            year = int(year_s)
            if year < 100:
                year += 2000
            try:
                return datetime.date(year, mon, int(day_s)).isoformat()
            except ValueError:
                pass  # mis. tanggal 31 di bulan yg cuma 30 hari -- lanjut ke error di bawah
    if val.isdigit() and 1 <= int(val) <= 100000:
        return (_EXCEL_EPOCH + datetime.timedelta(days=int(val))).isoformat()
    raise ValueError(
        f"tanggal '{val}' tidak dikenali formatnya "
        f"(harus YYYY-MM-DD, DD/MM/YYYY, atau DD-Mon-YY)"
    )



def _parse_datetime(val):
    """Sama seperti _parse_date() tapi untuk kolom timestamp (ikut jam)."""
    for fmt in _DATETIME_FORMATS:
        try:
            return datetime.datetime.strptime(val, fmt).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            continue
    try:
        return _parse_date(val) + " 00:00:00"
    except ValueError:
        pass
    raise ValueError(
        f"tanggal/waktu '{val}' tidak dikenali formatnya "
        f"(harus YYYY-MM-DD HH:MM atau DD/MM/YYYY HH:MM)"
    )



def insert_rows(rows):
    """rows: list of dict {field_key: value} (persis format `r` yang
    dipakai renderSqlOutput() di hem.html). Key di luar HEM_FIELD_KEYS
    diabaikan diam-diam (bukan error), supaya aman kalau frontend nanti
    kirim field ekstra.

    Return: (jumlah_baris_berhasil_masuk, list_pesan_error_per_baris).
    Baris yang gagal di-cast (mis. kolom angka diisi teks) DILEWATI
    (tidak membatalkan baris lain), errornya dilaporkan per baris."""
    if not rows:
        return 0, ["Tidak ada baris data yang dikirim."]

    used_keys = [
        k for k in HEM_FIELD_KEYS
        if any(str(r.get(k, "") or "").strip() != "" for r in rows)
    ]
    if not used_keys:
        return 0, ["Tidak ada kolom terisi di baris manapun."]

    values = []
    errors = []
    for i, r in enumerate(rows):
        try:
            row_values = []
            for k in used_keys:
                try:
                    row_values.append(_cast_value(HEM_FIELD_TYPE[k], r.get(k), key=k))
                except ValueError as e:
                    raise ValueError(f"kolom '{k}': {e}") from None
            values.append(row_values)
        except ValueError as e:
            errors.append(f"Baris {i + 1}: {e}")


    if not values:
        return 0, errors

    cols_sql = ", ".join(f'"{k}"' for k in used_keys)
    query = f'INSERT INTO {TABLE_NAME} ({cols_sql}) VALUES %s'

    conn = get_connection()
    try:
        with conn:
            with conn.cursor() as cur:
                psycopg2.extras.execute_values(cur, query, values)
    finally:
        conn.close()

    return len(values), errors


def build_create_table_sql():
    """DDL bantuan kalau tabel data_semesta BELUM ada di Postgres-nya --
    ditampilkan lewat /debug/hem-schema buat di-copy-paste manual ke tab
    Query Railway. Sengaja TIDAK dieksekusi otomatis oleh aplikasi (biar
    perubahan schema tetap sepenuhnya keputusan/kontrol Anda)."""
    type_sql = {
        "number": "NUMERIC", "text": "TEXT", "textarea": "TEXT",
        "date": "DATE", "datetime": "TIMESTAMP",
    }
    cols = ",\n  ".join(f'"{k}" {type_sql[t]}' for k, t in HEM_FIELDS)
    return f'CREATE TABLE IF NOT EXISTS {TABLE_NAME} (\n  id SERIAL PRIMARY KEY,\n  {cols}\n);'